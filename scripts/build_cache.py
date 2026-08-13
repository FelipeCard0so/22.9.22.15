import json
import os
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(__file__).resolve().parent / "data"
FILES = {tech: ROOT / f"{tech}_BASE.xlsx" for tech in ("2G", "3G", "4G", "5G")}
FIELDS = {
    "site": "[P]SITE", "uf": "[P]UF", "band": "[P]BANDA_OPERACAO",
    "azimuth": "[P]AZIMUTH", "bcch": "[P]BCCH", "psc": "[P]PSC",
    "pci": "[P]PCI", "bandwidth": "[P]BANDWIDTH", "city": "[P]CIDADE",
    "district": "[P]BAIRRO", "address": "[P]ENDERECO", "lat": "[P]LATITUDE",
    "lon": "[P]LONGITUDE", "mimo": "[P]MIMO", "dl_earfcn": "[P]DL_EARFCN",
    "dl_uarfcn": "[P]DL_UARFCN",
}


def value(row, indexes, field):
    index = indexes.get(FIELDS[field])
    if index is None or index >= len(row) or row[index] is None:
        return ""
    text = str(row[index]).strip()
    return "" if text.lower() == "nan" else text


def build():
    OUT.mkdir(exist_ok=True)
    for tech, path in FILES.items():
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(item).strip() if item is not None else "" for item in next(iterator)]
        indexes = {name: index for index, name in enumerate(headers)}
        records = []
        for row in iterator:
            site = value(row, indexes, "site")
            if not site:
                continue
            record = {field: value(row, indexes, field) for field in FIELDS}
            record["tech"] = tech
            records.append(record)
        workbook.close()
        target = OUT / f"{tech}.json"
        target.write_text(json.dumps(records, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        print(f"{tech}: {len(records):,} registros -> {target.stat().st_size / 1024 / 1024:.2f} MB")


if __name__ == "__main__":
    build()
